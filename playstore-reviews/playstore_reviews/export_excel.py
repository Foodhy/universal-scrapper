from __future__ import annotations

from datetime import datetime, timezone
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from playstore_reviews.config import Settings
from playstore_reviews.models import ScrapeResult

HEADERS = [
    "app_id",
    "pais",
    "idioma",
    "dispositivo",
    "estrellas",
    "fecha_utc",
    "usuario",
    "comentario",
    "version_app",
    "votos_util",
    "respuesta_desarrollador",
    "fecha_respuesta_utc",
    "review_id",
    "url_resena",
    "avatar",
    "orden",
    "estrategia",
    "extraido_utc",
]


def _style_header(sheet) -> None:
    fill = PatternFill("solid", fgColor="1F4E3D")
    font = Font(color="FFFFFF", bold=True)
    for cell in sheet[1]:
        cell.fill = fill
        cell.font = font
        cell.alignment = Alignment(vertical="center", wrap_text=True)
    sheet.auto_filter.ref = sheet.dimensions
    sheet.freeze_panes = "A2"
    sheet.row_dimensions[1].height = 22


def write_workbook(result: ScrapeResult, settings: Settings) -> Path:
    settings.output_dir.mkdir(parents=True, exist_ok=True)
    stamp = datetime.now(timezone.utc).strftime("%Y%m%d-%H%M%S")
    path = settings.output_dir / (
        f"reviews-{settings.app_id}-{settings.country}-{stamp}.xlsx"
    )

    book = Workbook()
    reviews = book.active
    reviews.title = "resenas"
    reviews.append(HEADERS)
    for review in result.reviews:
        reviews.append(
            [
                review.app_id,
                review.country,
                review.lang,
                review.device,
                review.score,
                review.review_date,
                review.user_name,
                review.content,
                review.app_version,
                review.thumbs_up,
                review.developer_reply,
                review.developer_reply_date,
                review.review_id,
                review.review_url,
                review.user_image,
                review.sort,
                review.strategy,
                review.fetched_at,
            ]
        )
    _style_header(reviews)
    for row in reviews.iter_rows(min_row=2, max_col=len(HEADERS)):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    widths = {
        1: 28,
        6: 22,
        7: 24,
        8: 70,
        11: 40,
        13: 36,
        14: 42,
        15: 28,
    }
    for index in range(1, len(HEADERS) + 1):
        reviews.column_dimensions[get_column_letter(index)].width = widths.get(index, 16)

    summary = book.create_sheet("resumen")
    summary.append(["campo", "valor"])
    summary.append(["app_id", settings.app_id])
    summary.append(["pais", settings.country])
    summary.append(["idioma", settings.lang])
    summary.append(["orden", settings.sort])
    summary.append(["estrategia_pedida", settings.strategy])
    summary.append(["limite_por_consulta", settings.per_query_limit])
    summary.append(["proxies", len(settings.proxies)])
    summary.append(["reseñas", len(result.reviews)])
    summary.append(["paises_en_filas", ", ".join(sorted({review.country for review in result.reviews}))])
    summary.append(["detenido", "si" if result.stopped else "no"])
    summary.append([])
    summary.append(["dispositivo", "estrellas", "obtenidas", "estrategia", "error"])
    for query in result.queries:
        summary.append([query.device, query.score, query.fetched, query.strategy, query.error or ""])
    _style_header(summary)
    summary.column_dimensions["A"].width = 24
    summary.column_dimensions["B"].width = 22
    summary.column_dimensions["E"].width = 60

    book.save(path)
    return path
