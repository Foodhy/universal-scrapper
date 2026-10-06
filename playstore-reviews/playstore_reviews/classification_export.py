from __future__ import annotations

from datetime import datetime, timezone
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from playstore_reviews.insights import apply_threshold


def write_classification_excel(
    reviews: list[dict],
    problems: list[dict],
    threshold: float,
    output_dir: Path,
) -> Path:
    output_dir.mkdir(parents=True, exist_ok=True)
    stamp = datetime.now(timezone.utc).strftime("%Y%m%d-%H%M%S")
    path = output_dir / f"clasificacion-{stamp}.xlsx"
    book = Workbook()
    sheet = book.active
    sheet.title = "clasificacion"
    headers = [
        "fecha_utc",
        "pais",
        "dispositivo",
        "estrellas",
        "usuario",
        "comentario",
        "problemas",
    ]
    for problem in problems:
        headers.append(f"{problem['id']}_prob")
    headers.extend(["review_id", "app_id", "version_app", "umbral"])
    sheet.append(headers)
    for review in reviews:
        scored = apply_threshold(review.get("answers") or {}, threshold)
        matched = [
            problem["title"]
            for problem in problems
            if scored.get(problem["id"], {}).get("matched")
        ]
        row = [
            review.get("review_date"),
            review.get("country"),
            review.get("device"),
            review.get("score"),
            review.get("user_name"),
            review.get("content"),
            ", ".join(matched),
        ]
        for problem in problems:
            row.append(scored.get(problem["id"], {}).get("noul"))
        row.extend([review.get("review_id"), review.get("app_id"), review.get("app_version"), threshold])
        sheet.append(row)
    fill = PatternFill("solid", fgColor="1F4E3D")
    font = Font(color="FFFFFF", bold=True)
    for cell in sheet[1]:
        cell.fill = fill
        cell.font = font
        cell.alignment = Alignment(wrap_text=True, vertical="center")
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    for index in range(1, len(headers) + 1):
        sheet.column_dimensions[get_column_letter(index)].width = 18 if index != 6 else 70
    book.save(path)
    return path
